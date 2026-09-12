from fastapi.testclient import TestClient


def test_get_setup_full(client: TestClient):
    resp = client.get("/api/v1/setup/2026")
    assert resp.status_code == 200
    body = resp.json()
    assert "Roundesk" in body["income_sources"]
    assert "Food" in body["expense_categories"]
    assert {"category": "Food", "subcategory": "Chicken"} in body["subcategories"]
    assert "Entertainment" in body["free_money_categories"]
    assert "Education" in body["investment_areas"]
    assert len(body["recurring_expenses"]) == 6


def test_subcategories_and_classifications_are_clean(client: TestClient):
    # tbl_Subcategories originally overlapped tbl_Classifications (E15:E19)
    # and tbl_Recurring (A24:F30) in the source workbook, leaking their
    # content as fake (category, subcategory) pairs, before being relocated
    # to a dedicated H:I range (see scripts/relocate_subcategories_table.py).
    # This is a regression guard for that fix.
    body = client.get("/api/v1/setup/2026").json()
    categories_seen = {pair["category"] for pair in body["subcategories"]}
    assert "Classification" not in categories_seen
    assert not any("Necessary" in c or "Unnecessary" in c for c in categories_seen)
    assert set(body["classifications"]) == {
        "Necessary — Planned", "Necessary — Unplanned",
        "Unnecessary — Planned", "Unnecessary — Unplanned",
    }
    # Previously unreadable due to the overlap - recovered per spec section 7.
    assert {"category": "Food", "subcategory": "Fast Food"} in body["subcategories"]
    assert {"category": "Food", "subcategory": "Snacks"} in body["subcategories"]
    assert {"category": "Motorcycle", "subcategory": "Fuel"} in body["subcategories"]
    assert {"category": "Motorcycle", "subcategory": "Oil"} in body["subcategories"]
    assert {"category": "Motorcycle", "subcategory": "Maintenance"} in body["subcategories"]


def test_recurring_expense_monthly_reserve_computed(client: TestClient):
    resp = client.get("/api/v1/setup/2026/recurring-expenses")
    assert resp.status_code == 200
    by_expense = {r["expense"]: r for r in resp.json()}
    # Expected amounts are all 0 in the source workbook, so reserves are 0,
    # but the frequency-based math itself must run without error.
    assert by_expense["Rent"]["frequency"] == "Monthly"
    assert by_expense["Insurance"]["frequency"] == "Yearly"


def test_get_individual_setup_lists(client: TestClient):
    assert "Housing" in client.get("/api/v1/setup/2026/categories").json()
    assert "Family Support" in client.get("/api/v1/setup/2026/income-sources").json()
    assert "Career" in client.get("/api/v1/setup/2026/investment-areas").json()
    assert "Shopping" in client.get("/api/v1/setup/2026/free-money-categories").json()


def test_add_subcategory_succeeds_with_spare_row(client: TestClient):
    # tbl_Subcategories was relocated to H7:I45 with real spare capacity
    # (20 blank rows) specifically so this works without any workaround.
    resp = client.post(
        "/api/v1/setup/2026/subcategories", json={"category": "Food", "subcategory": "Yogurt"}
    )
    assert resp.status_code == 201
    assert resp.json() == {"category": "Food", "subcategory": "Yogurt"}

    subs = client.get("/api/v1/setup/2026/subcategories").json()
    assert {"category": "Food", "subcategory": "Yogurt"} in subs

    # Now usable on a transaction.
    tx = client.post(
        "/api/v1/months/2026/9/transactions",
        json={
            "date": "2026-09-11", "type": "Expense", "category": "Food",
            "subcategory": "Yogurt", "item": "Greek yogurt", "amount": 5,
        },
    )
    assert tx.status_code == 201


def test_add_subcategory_rejected_when_table_full(client: TestClient, data_dir):
    import openpyxl

    path = data_dir / "Personal_Finance_2026_V1.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb["SETUP"]
    # Fill every spare row (26-45) so the table genuinely has no room left,
    # to exercise the "block full" rule itself rather than relying on
    # whatever the real workbook's current capacity happens to be.
    for row in range(26, 46):
        ws.cell(row=row, column=8, value="Food")
        ws.cell(row=row, column=9, value=f"Filler{row}")
    wb.save(path)

    resp = client.post(
        "/api/v1/setup/2026/subcategories", json={"category": "Food", "subcategory": "Yogurt"}
    )
    assert resp.status_code == 409
    assert resp.json()["detail"]["error"] == "table_full"


def test_add_subcategory_unknown_category_rejected(client: TestClient):
    resp = client.post(
        "/api/v1/setup/2026/subcategories", json={"category": "NotACategory", "subcategory": "X"}
    )
    assert resp.status_code == 400
    assert resp.json()["detail"]["error"] == "invalid_category"


def test_add_duplicate_subcategory_is_idempotent(client: TestClient):
    first = client.post(
        "/api/v1/setup/2026/subcategories", json={"category": "Food", "subcategory": "Chicken"}
    )
    assert first.status_code == 201
    subs = client.get("/api/v1/setup/2026/subcategories").json()
    chicken_count = sum(1 for s in subs if s == {"category": "Food", "subcategory": "Chicken"})
    assert chicken_count == 1


def test_add_free_money_category_with_spare_row(client: TestClient, data_dir):
    import openpyxl

    # tbl_FreeMoneyCategories is genuinely at capacity in the real fixture
    # (5/5 rows filled) - free one up to test the actual add path, same
    # situation as tbl_Subcategories before it was relocated.
    path = data_dir / "Personal_Finance_2026_V1.xlsx"
    wb = openpyxl.load_workbook(path)
    wb["SETUP"]["A20"] = None  # last row of tbl_FreeMoneyCategories ("Other")
    wb.save(path)

    add = client.post("/api/v1/setup/2026/free-money-categories", json={"value": "Gaming"})
    assert add.status_code == 201
    assert "Gaming" in client.get("/api/v1/setup/2026/free-money-categories").json()


def test_edit_and_delete_free_money_category(client: TestClient):
    # Uses an existing entry rather than adding a new one, since the real
    # list has no spare capacity.
    rename = client.put(
        "/api/v1/setup/2026/free-money-categories/Other", json={"value": "Miscellaneous"}
    )
    assert rename.status_code == 200
    categories = client.get("/api/v1/setup/2026/free-money-categories").json()
    assert "Miscellaneous" in categories
    assert "Other" not in categories

    delete = client.delete("/api/v1/setup/2026/free-money-categories/Miscellaneous")
    assert delete.status_code == 204
    assert "Miscellaneous" not in client.get("/api/v1/setup/2026/free-money-categories").json()


def test_delete_unknown_list_value_returns_404(client: TestClient):
    resp = client.delete("/api/v1/setup/2026/investment-areas/NotARealArea")
    assert resp.status_code == 404
    assert resp.json()["detail"]["error"] == "not_found"


def test_add_duplicate_list_value_is_idempotent(client: TestClient):
    first = client.post("/api/v1/setup/2026/income-sources", json={"value": "Roundesk"})
    assert first.status_code == 201
    sources = client.get("/api/v1/setup/2026/income-sources").json()
    assert sources.count("Roundesk") == 1


def test_edit_and_delete_subcategory(client: TestClient):
    add = client.post(
        "/api/v1/setup/2026/subcategories", json={"category": "Housing", "subcategory": "Repairs"}
    )
    assert add.status_code == 201

    rename = client.put(
        "/api/v1/setup/2026/subcategories/Housing/Repairs", json={"subcategory": "Maintenance"}
    )
    assert rename.status_code == 200
    subs = client.get("/api/v1/setup/2026/subcategories").json()
    assert {"category": "Housing", "subcategory": "Maintenance"} in subs
    assert {"category": "Housing", "subcategory": "Repairs"} not in subs

    delete = client.delete("/api/v1/setup/2026/subcategories/Housing/Maintenance")
    assert delete.status_code == 204
    subs = client.get("/api/v1/setup/2026/subcategories").json()
    assert {"category": "Housing", "subcategory": "Maintenance"} not in subs


def test_invalid_list_kind_rejected(client: TestClient):
    resp = client.post("/api/v1/setup/2026/not-a-real-list", json={"value": "X"})
    assert resp.status_code == 422
