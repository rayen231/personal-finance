import openpyxl
from openpyxl.utils.cell import range_boundaries
from fastapi.testclient import TestClient


def test_missing_year_is_created_from_template_and_inherits_setup(client: TestClient, data_dir):
    assert not (data_dir / "Personal_Finance_2027_V1.xlsx").exists()

    resp = client.post(
        "/api/v1/months/2027/1/transactions",
        json={
            "date": "2027-01-05",
            "type": "Expense",
            "category": "Food",
            "subcategory": "Chicken",
            "item": "New year test",
            "amount": 22,
        },
    )
    assert resp.status_code == 201

    new_file = data_dir / "Personal_Finance_2027_V1.xlsx"
    assert new_file.exists()

    wb = openpyxl.load_workbook(new_file)
    ws = wb["JANUARY"]
    assert ws["B2"].value == "Active"  # auto-activated on first write
    assert ws.cell(row=43, column=3).value == "Food"

    setup = wb["SETUP"]
    min_col, min_row, max_col, max_row = range_boundaries(setup.tables["tbl_ExpenseCategories"].ref)
    categories = [setup.cell(row=r, column=min_col).value for r in range(min_row + 1, max_row + 1)]
    assert "Food" in categories
    assert "Housing" in categories
