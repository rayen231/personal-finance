"""One-time migration: add a hidden TransactionID column to every month's
transaction table (tbl_<MONTH>_Transactions), so API-created rows have a
stable key that survives updates/deletes elsewhere in the table.

Usage:
    python scripts/migrate_add_transaction_id.py <path-to-workbook.xlsx>

Modifies the file in place. Run it once per workbook (idempotent: skips
months that already have the column).
"""

import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import TableColumn

MONTHS = [
    "JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
    "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER",
]

ID_COLUMN_NAME = "TransactionID"


def migrate(path: Path) -> None:
    wb = openpyxl.load_workbook(path)

    for month in MONTHS:
        if month not in wb.sheetnames:
            continue
        ws = wb[month]
        table_name = f"tbl_{month}_Transactions"
        if table_name not in ws.tables:
            print(f"  [skip] {month}: table {table_name} not found")
            continue

        table = ws.tables[table_name]
        existing_cols = [c.name for c in table.tableColumns]
        if ID_COLUMN_NAME in existing_cols:
            print(f"  [ok]   {month}: already migrated")
            continue

        min_col, min_row, max_col, max_row = openpyxl.utils.cell.range_boundaries(table.ref)
        new_col_idx = max_col + 1
        new_col_letter = get_column_letter(new_col_idx)

        header_cell = ws.cell(row=min_row, column=new_col_idx, value=ID_COLUMN_NAME)

        next_col_id = max((c.id for c in table.tableColumns), default=0) + 1
        table.tableColumns.append(TableColumn(id=next_col_id, name=ID_COLUMN_NAME))
        table.ref = f"{get_column_letter(min_col)}{min_row}:{new_col_letter}{max_row}"

        ws.column_dimensions[new_col_letter].hidden = True

        print(f"  [done] {month}: added column {new_col_letter} ('{ID_COLUMN_NAME}')")

    wb.save(path)
    print(f"Saved: {path}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(1)
    migrate(Path(sys.argv[1]))
