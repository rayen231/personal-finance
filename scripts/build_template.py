"""Builds the blank template workbook used to create a new year's file.

Starts from the current (migrated) workbook and clears everything that's
actually *data* - transactions, actual income, plans, investments - while
keeping every formula, table, dropdown and the SETUP defaults intact.

Usage:
    python scripts/build_template.py <source.xlsx> <template_output.xlsx>
"""

import sys
from pathlib import Path

import openpyxl

MONTHS = [
    "JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE",
    "JULY", "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER",
]


def clear_range(ws, min_row, max_row, min_col, max_col):
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            ws.cell(row=row, column=col).value = None


def build_template(source_path: Path, output_path: Path) -> None:
    wb = openpyxl.load_workbook(source_path)

    for month in MONTHS:
        if month not in wb.sheetnames:
            continue
        ws = wb[month]

        ws["B2"] = "Not Started"  # month status

        # Income planning: actual (col C) + expected (col B) for source rows.
        clear_range(ws, 6, 9, 2, 3)

        # Necessary expenses planning: planned (col B) for category rows.
        clear_range(ws, 14, 20, 2, 2)

        # Free money planned allocation (col G).
        clear_range(ws, 22, 26, 7, 7)

        # Investments planning block (Date/Area/Subcategory/Amount/Notes).
        clear_range(ws, 25, 34, 1, 6)

        # Main transactions table: all data columns (A-F, J, K, L); leave
        # the formula columns G/H/I (=B, =C, =D) untouched.
        table_name = f"tbl_{month}_Transactions"
        if table_name in ws.tables:
            min_col, min_row, max_col, max_row = openpyxl.utils.cell.range_boundaries(
                ws.tables[table_name].ref
            )
            data_row_start = min_row + 1
            for row in range(data_row_start, max_row + 1):
                for col in (1, 2, 3, 4, 5, 6, 10, 11, 12):
                    ws.cell(row=row, column=col).value = None

    wb.save(output_path)
    print(f"Template written to {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    build_template(Path(sys.argv[1]), Path(sys.argv[2]))
