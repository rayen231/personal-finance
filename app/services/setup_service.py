"""Reads configuration lists (categories, subcategories, etc.) from SETUP."""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.workbook import Workbook

from app.services.excel_service import LabelNotFoundError, TableFullError, WorkbookError

VALID_TX_TYPES = ["Expense", "Free Money", "Investment"]

# Mirrors the IF() formula in tbl_Recurring's "Monthly Reserve" column, so
# the API can return a real number instead of an unevaluated formula string
# (openpyxl never computes formulas).
_FREQUENCY_DIVISORS = {
    "Monthly": 1,
    "Yearly": 12,
    "Every 6 months": 6,
    "Every 3 months": 3,
}


def _monthly_reserve(frequency: str, expected_amount: float) -> float:
    if not expected_amount:
        return 0.0
    if frequency == "Weekly":
        return expected_amount * 52 / 12
    return expected_amount / _FREQUENCY_DIVISORS.get(frequency, 1)


def _table_single_column(ws, table_name: str) -> list[str]:
    if table_name not in ws.tables:
        raise WorkbookError(f"Table {table_name} not found in SETUP")
    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)
    values = []
    for row in range(min_row + 1, max_row + 1):
        val = ws.cell(row=row, column=min_col).value
        if val is not None:
            values.append(str(val))
    return values


def _rows_claimed_by_other_tables(ws, own_table_name: str, min_col: int, max_col: int) -> set[int]:
    """SETUP has at least one known case (tbl_Subcategories vs tbl_Classifications
    and tbl_Recurring) where a table's declared ref genuinely overlaps another
    table's real range in the source workbook - not a blank/unused area. Reading
    or writing those rows would silently corrupt whichever table actually owns
    them, so every row/col range is checked against all *other* tables on the
    sheet and excluded if it overlaps, rather than trusting one table's own ref."""
    excluded: set[int] = set()
    for name in ws.tables:
        if name == own_table_name:
            continue
        o_min_col, o_min_row, o_max_col, o_max_row = range_boundaries(ws.tables[name].ref)
        if o_max_col < min_col or o_min_col > max_col:
            continue  # no column overlap
        excluded.update(range(o_min_row, o_max_row + 1))
    return excluded


def _table_two_columns(ws, table_name: str) -> list[tuple[str, str]]:
    if table_name not in ws.tables:
        raise WorkbookError(f"Table {table_name} not found in SETUP")
    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)
    excluded_rows = _rows_claimed_by_other_tables(ws, table_name, min_col, max_col)
    pairs = []
    for row in range(min_row + 1, max_row + 1):
        if row in excluded_rows:
            continue
        a = ws.cell(row=row, column=min_col).value
        b = ws.cell(row=row, column=min_col + 1).value
        if a is not None and b is not None:
            pairs.append((str(a), str(b)))
    return pairs


@dataclass
class SetupConfig:
    income_sources: list[str] = field(default_factory=list)
    expense_categories: list[str] = field(default_factory=list)
    subcategories: list[tuple[str, str]] = field(default_factory=list)  # (category, subcategory)
    free_money_categories: list[str] = field(default_factory=list)
    investment_areas: list[str] = field(default_factory=list)
    classifications: list[str] = field(default_factory=list)

    def subcategories_for(self, category: str) -> list[str]:
        return [sub for cat, sub in self.subcategories if cat == category]

    def is_valid_category_for_type(self, tx_type: str, category: str) -> bool:
        if tx_type == "Expense":
            return category in self.expense_categories
        if tx_type == "Free Money":
            return category in self.free_money_categories
        if tx_type == "Investment":
            return category in self.investment_areas
        return False

    def is_valid_subcategory(self, tx_type: str, category: str, subcategory: str) -> bool:
        if tx_type == "Expense":
            return subcategory in self.subcategories_for(category)
        # Free Money / Investment categories don't have a SETUP subcategory table;
        # subcategory is a free-text label for those types.
        return True


def read_setup_config(wb: Workbook) -> SetupConfig:
    ws = wb["SETUP"]
    expense_categories = _table_single_column(ws, "tbl_ExpenseCategories")
    raw_subcategories = _table_two_columns(ws, "tbl_Subcategories")
    # Belt-and-suspenders beyond the overlap check above: a row can hold a
    # stray label in the category cell without being claimed by any other
    # table (e.g. a section title sitting one row above a table it isn't
    # part of). Rather than guess what such a row "should" say, drop any
    # pair whose category isn't a real, known expense category.
    subcategories = [(c, s) for c, s in raw_subcategories if c in expense_categories]
    return SetupConfig(
        income_sources=_table_single_column(ws, "tbl_IncomeSources"),
        expense_categories=expense_categories,
        subcategories=subcategories,
        free_money_categories=_table_single_column(ws, "tbl_FreeMoneyCategories"),
        investment_areas=_table_single_column(ws, "tbl_InvestmentAreas"),
        classifications=_table_single_column(ws, "tbl_Classifications"),
    )


def read_recurring_expenses(wb: Workbook) -> list[dict]:
    ws = wb["SETUP"]
    table_name = "tbl_Recurring"
    if table_name not in ws.tables:
        raise WorkbookError(f"Table {table_name} not found in SETUP")
    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)

    results = []
    for row in range(min_row + 1, max_row + 1):
        expense = ws.cell(row=row, column=min_col).value
        if expense is None:
            continue
        category = ws.cell(row=row, column=min_col + 1).value
        frequency = ws.cell(row=row, column=min_col + 2).value
        expected_amount = ws.cell(row=row, column=min_col + 3).value or 0
        notes = ws.cell(row=row, column=min_col + 5).value
        results.append({
            "expense": expense,
            "category": category,
            "frequency": frequency,
            "expected_amount": expected_amount,
            "monthly_reserve": _monthly_reserve(frequency, expected_amount),
            "notes": notes,
        })
    return results


def _find_value_row(ws, table_name: str, value: str) -> int:
    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)
    excluded_rows = _rows_claimed_by_other_tables(ws, table_name, min_col, max_col)
    for row in range(min_row + 1, max_row + 1):
        if row in excluded_rows:
            continue
        if ws.cell(row=row, column=min_col).value == value:
            return row
    raise LabelNotFoundError(f"'{value}' not found in {table_name}.")


def add_single_value(wb: Workbook, table_name: str, value: str) -> bool:
    """Generic add for a single-column SETUP list (categories, income
    sources, etc.). Returns True if written, False if it already existed
    (idempotent no-op)."""
    ws = wb["SETUP"]
    if table_name not in ws.tables:
        raise WorkbookError(f"Table {table_name} not found in SETUP")

    try:
        _find_value_row(ws, table_name, value)
        return False  # already exists
    except LabelNotFoundError:
        pass

    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)
    excluded_rows = _rows_claimed_by_other_tables(ws, table_name, min_col, max_col)
    for row in range(min_row + 1, max_row + 1):
        if row in excluded_rows:
            continue
        if ws.cell(row=row, column=min_col).value is None:
            ws.cell(row=row, column=min_col, value=value)
            return True

    raise TableFullError(
        f"{table_name} is full (no blank row up to row {max_row}). Add rows to the table in Excel first."
    )


def edit_single_value(wb: Workbook, table_name: str, old_value: str, new_value: str) -> None:
    """Renames an entry in a single-column SETUP list in place. Does not
    cascade to historical transactions that reference the old value - they
    keep referencing whatever text they already stored."""
    ws = wb["SETUP"]
    if table_name not in ws.tables:
        raise WorkbookError(f"Table {table_name} not found in SETUP")
    row = _find_value_row(ws, table_name, old_value)
    min_col, _, _, _ = range_boundaries(ws.tables[table_name].ref)
    ws.cell(row=row, column=min_col, value=new_value)


def delete_single_value(wb: Workbook, table_name: str, value: str) -> None:
    """Clears an entry from a single-column SETUP list. Does not touch any
    historical transactions already using this value."""
    ws = wb["SETUP"]
    if table_name not in ws.tables:
        raise WorkbookError(f"Table {table_name} not found in SETUP")
    row = _find_value_row(ws, table_name, value)
    min_col, _, _, _ = range_boundaries(ws.tables[table_name].ref)
    ws.cell(row=row, column=min_col).value = None


def _find_subcategory_row(ws, category: str, subcategory: str) -> int:
    table_name = "tbl_Subcategories"
    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)
    excluded_rows = _rows_claimed_by_other_tables(ws, table_name, min_col, max_col)
    for row in range(min_row + 1, max_row + 1):
        if row in excluded_rows:
            continue
        if (
            ws.cell(row=row, column=min_col).value == category
            and ws.cell(row=row, column=min_col + 1).value == subcategory
        ):
            return row
    raise LabelNotFoundError(f"Subcategory '{subcategory}' not found under category '{category}'.")


def edit_subcategory(wb: Workbook, category: str, old_subcategory: str, new_subcategory: str) -> None:
    ws = wb["SETUP"]
    row = _find_subcategory_row(ws, category, old_subcategory)
    min_col, _, _, _ = range_boundaries(ws.tables["tbl_Subcategories"].ref)
    ws.cell(row=row, column=min_col + 1, value=new_subcategory)


def delete_subcategory(wb: Workbook, category: str, subcategory: str) -> None:
    ws = wb["SETUP"]
    row = _find_subcategory_row(ws, category, subcategory)
    min_col, _, _, _ = range_boundaries(ws.tables["tbl_Subcategories"].ref)
    ws.cell(row=row, column=min_col).value = None
    ws.cell(row=row, column=min_col + 1).value = None


def add_subcategory(wb: Workbook, category: str, subcategory: str) -> bool:
    """Returns True if a new row was written, False if the pair already
    existed (idempotent no-op) - callers use this to decide whether the
    workbook actually needs to be re-saved."""
    ws = wb["SETUP"]
    table_name = "tbl_Subcategories"
    if table_name not in ws.tables:
        raise WorkbookError(f"Table {table_name} not found in SETUP")

    config = read_setup_config(wb)
    if category not in config.expense_categories:
        raise LabelNotFoundError(f"'{category}' is not a known expense category.")
    if subcategory in config.subcategories_for(category):
        return False  # already exists - adding is idempotent

    min_col, min_row, max_col, max_row = range_boundaries(ws.tables[table_name].ref)
    excluded_rows = _rows_claimed_by_other_tables(ws, table_name, min_col, max_col)
    for row in range(min_row + 1, max_row + 1):
        if row in excluded_rows:
            continue
        if ws.cell(row=row, column=min_col).value is None:
            ws.cell(row=row, column=min_col, value=category)
            ws.cell(row=row, column=min_col + 1, value=subcategory)
            return True

    raise TableFullError(
        f"{table_name} is full (no blank row up to row {max_row}). Add rows to the table in Excel first."
    )
