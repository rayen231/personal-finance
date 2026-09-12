"""Ties storage + locking + per-year workbook lifecycle together.

Business/API code should go through this rather than touching
StorageService or openpyxl directly - this is the seam where Excel could
later be swapped for a real database.
"""

from __future__ import annotations

import shutil
from contextlib import contextmanager
from threading import Lock

from openpyxl.workbook.workbook import Workbook

from app.config import Settings, get_settings
from app.services import excel_service, setup_service
from app.services.storage_service import StorageService, WorkbookNotFoundError, get_storage_service
from app.services.sync_state import SyncStateStore

_year_locks: dict[int, Lock] = {}


def _lock_for(year: int) -> Lock:
    return _year_locks.setdefault(year, Lock())


class WriteHandle:
    """Wraps a Workbook opened via open_for_write. `changed` defaults to True
    (most writers unconditionally modify something); callers that might do
    nothing - an idempotent retry, a no-op update, a fully-failed batch -
    should set it False so a no-op doesn't upload a new blob or bump the
    version (which a mobile client could otherwise mistake for new data)."""

    def __init__(self, wb: Workbook):
        self.wb = wb
        self.changed = True


class WorkbookRepository:
    def __init__(self, storage: StorageService, settings: Settings, sync_state: SyncStateStore):
        self.storage = storage
        self.settings = settings
        self.sync_state = sync_state

    def _filename(self, year: int) -> str:
        return self.settings.workbook_filename(year)

    def _ensure_workbook_exists(self, year: int) -> None:
        filename = self._filename(year)
        if self.storage.exists(filename):
            return

        # Lazily create from the template, inheriting SETUP from the latest
        # existing year so user-added categories/subcategories carry forward.
        template_name = self.settings.template_filename
        if not self.storage.exists(template_name):
            raise WorkbookNotFoundError(
                f"Neither {filename} nor template {template_name} exist. "
                "Create a template workbook first."
            )

        new_content = self.storage.download(template_name)
        latest_year = self._find_latest_existing_year(before=year)
        if latest_year is not None:
            new_content = self._clone_setup_from(latest_year, new_content)

        self.storage.upload(filename, new_content)

    def _find_latest_existing_year(self, before: int) -> int | None:
        candidates = [y for y in range(before - 1, before - 20, -1) if self.storage.exists(self._filename(y))]
        return candidates[0] if candidates else None

    def _clone_setup_from(self, source_year: int, template_content: bytes) -> bytes:
        source_wb = excel_service.load_workbook(self.storage.download(self._filename(source_year)))
        target_wb = excel_service.load_workbook(template_content)

        source_ws, target_ws = source_wb["SETUP"], target_wb["SETUP"]
        for table_name in source_ws.tables:
            if table_name not in target_ws.tables:
                continue
            src_bounds = excel_service.range_boundaries(source_ws.tables[table_name].ref)
            for row in range(src_bounds[1] + 1, src_bounds[3] + 1):
                for col in range(src_bounds[0], src_bounds[2] + 1):
                    value = source_ws.cell(row=row, column=col).value
                    target_ws.cell(row=row, column=col, value=value)

        return excel_service.save_workbook(target_wb)

    @contextmanager
    def open_for_write(self, year: int):
        """Yields a WriteHandle wrapping a loaded Workbook; saves + uploads
        it and bumps server_version on clean exit, unless the caller cleared
        handle.changed to signal nothing actually happened."""
        with _lock_for(year):
            self._ensure_workbook_exists(year)
            content = self.storage.download(self._filename(year))
            handle = WriteHandle(excel_service.load_workbook(content))
            yield handle
            if not handle.changed:
                return
            new_content = excel_service.save_workbook(handle.wb)
            # Verify the saved bytes can be reopened before publishing them.
            excel_service.load_workbook(new_content)
            self.storage.upload(self._filename(year), new_content)
            self.sync_state.bump_version(year)

    def open_for_read(self, year: int) -> Workbook:
        with _lock_for(year):
            self._ensure_workbook_exists(year)
            content = self.storage.download(self._filename(year))
        return excel_service.load_workbook(content)

    def workbook_exists(self, year: int) -> bool:
        """Checks storage without lazily creating anything - used by callers
        (e.g. recurring-value carry-forward) that need to look at an earlier
        year but must never conjure a blank workbook into existence just by
        looking at it."""
        return self.storage.exists(self._filename(year))

    def get_setup(self, year: int) -> setup_service.SetupConfig:
        wb = self.open_for_read(year)
        return setup_service.read_setup_config(wb)


def get_workbook_repository() -> WorkbookRepository:
    settings = get_settings()
    storage = get_storage_service(settings)
    sync_state = SyncStateStore(storage)
    return WorkbookRepository(storage, settings, sync_state)
